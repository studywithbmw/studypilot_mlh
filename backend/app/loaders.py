"""Read PDFs, Word/PowerPoint files, plain text, web pages and YouTube transcripts,
then split them into overlapping chunks. Pure Python; no network except link import."""
import ipaddress
import re
import socket
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

from .config import settings

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}


class IngestError(Exception):
    """A problem with the file itself; the message is shown to the user."""


@dataclass
class Page:
    number: int | None  # 1-based page/slide for PDF/PPTX, None otherwise
    text: str
    label: str | None = None  # e.g. "03:20" for a video transcript window


@dataclass
class Chunk:
    text: str
    page: int | None
    index: int
    label: str | None = None


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------
def _read_pdf(path: Path) -> tuple[list[Page], int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            raise IngestError("This PDF is password protected.") from None
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(Page(i, text))
    if sum(len(p.text) for p in pages) < 50:
        raise IngestError("No readable text found. This looks like a scanned PDF (images only); run OCR on it first.")
    return pages, len(reader.pages)


def _read_docx(path: Path) -> list[Page]:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise IngestError("The Word file has no text.")
    return [Page(None, text)]


def _read_pptx(path: Path) -> list[Page]:
    try:
        from pptx import Presentation
    except ImportError:
        raise IngestError("PowerPoint support needs: pip install python-pptx") from None
    prs = Presentation(str(path))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            pages.append(Page(i, "\n".join(texts)))
    if not pages:
        raise IngestError("The presentation has no text.")
    return pages


def _read_text(path: Path) -> list[Page]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    if not text:
        raise IngestError("The file is empty.")
    return [Page(None, text)]


def load_pages(path: Path) -> tuple[list[Page], int | None]:
    """Return (pages with text, total page/slide count if the format has pages)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path), None
    if ext == ".pptx":
        pages = _read_pptx(path)
        return pages, max(p.number or 0 for p in pages)
    if ext in (".txt", ".md"):
        return _read_text(path), None
    raise IngestError(f"{ext or 'This'} files are not supported. Use PDF, DOCX, PPTX, TXT or MD.")


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------
_BREAKS = ("\n\n", "\n", ". ", "? ", "! ", "; ", ", ", " ")


def split_text(text: str, size: int, overlap: int) -> list[str]:
    """Greedy splitter: cut at the best natural boundary in the back half of the window."""
    text = text.strip()
    if len(text) <= size:
        return [text] if text else []
    overlap = min(overlap, size // 3)
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            window = text[start:end]
            for sep in _BREAKS:
                i = window.rfind(sep)
                if i >= size * 0.5:
                    end = start + i + len(sep)
                    break
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def split_pages(pages: list[Page]) -> list[Chunk]:
    chunks: list[Chunk] = []
    for page in pages:
        for piece in split_text(page.text, settings.chunk_size, settings.chunk_overlap):
            if len(piece) >= 20:
                chunks.append(Chunk(piece, page.number, len(chunks), page.label))
    return chunks


# ---------------------------------------------------------------------------
# Links (web page / YouTube)
# ---------------------------------------------------------------------------
def _is_public(address: str) -> bool:
    try:
        ip = ipaddress.ip_address(address)
    except ValueError:
        return False
    return not (ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified)


def check_public_url(url: str) -> str:
    """Reject anything that is not a public http(s) address (SSRF guard)."""
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise IngestError("Enter a full link starting with http:// or https://")
    host = parsed.hostname
    try:
        infos = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror:
        raise IngestError("That address could not be found.") from None
    if not infos or not all(_is_public(i[4][0]) for i in infos):
        raise IngestError("That link points to a private or internal address, which is not allowed.")
    return url.strip()


class _TextExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "nav", "footer", "header", "svg", "form", "aside"}
    _BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag in self._SKIP:
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag in self._SKIP and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):  # noqa: ANN001
        if self._in_title:
            self.title += data
        elif not self._skip and data.strip():
            self.parts.append(data)


def html_to_text(html: str) -> tuple[str, str]:
    p = _TextExtractor()
    p.feed(html)
    text = re.sub(r"[ \t]+", " ", "".join(p.parts))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return p.title.strip(), text


def fetch_web(url: str) -> tuple[str, list[Page]]:
    import httpx

    current = check_public_url(url)
    max_bytes = 5 * 1024 * 1024
    with httpx.Client(timeout=httpx.Timeout(10.0), follow_redirects=False, headers={"User-Agent": "StudyPilot/1.0"}) as client:
        for _ in range(5):
            try:
                resp = client.get(current)
            except httpx.HTTPError:
                raise IngestError("Could not download that page.") from None
            if resp.status_code in (301, 302, 303, 307, 308) and resp.headers.get("location"):
                current = check_public_url(urljoin(current, resp.headers["location"]))
                continue
            if resp.status_code >= 400:
                raise IngestError(f"The site answered with HTTP {resp.status_code}.")
            ctype = resp.headers.get("content-type", "")
            if "html" not in ctype and "text" not in ctype:
                raise IngestError("That link is not a web page (HTML or text).")
            title, text = html_to_text(resp.text[:max_bytes])
            if len(text) < 80:
                raise IngestError("Not enough readable text on that page.")
            return title or urlparse(url).hostname or "Web page", [Page(None, text)]
    raise IngestError("Too many redirects.")


def youtube_id(url: str) -> str | None:
    parsed = urlparse(url.strip())
    host = (parsed.hostname or "").lower()
    if host == "youtu.be":
        vid = parsed.path.lstrip("/").split("/")[0]
    elif host.endswith("youtube.com"):
        if parsed.path == "/watch":
            vid = (parse_qs(parsed.query).get("v") or [""])[0]
        else:
            m = re.match(r"^/(?:embed|shorts|live|v)/([A-Za-z0-9_-]{6,20})", parsed.path)
            vid = m.group(1) if m else ""
    else:
        return None
    return vid if re.fullmatch(r"[A-Za-z0-9_-]{6,20}", vid or "") else None


def _mmss(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"


def fetch_youtube(url: str) -> tuple[list[Page], int]:
    """Transcript grouped into ~60 second windows. Returns (pages, duration_seconds)."""
    vid = youtube_id(url)
    if not vid:
        raise IngestError("That does not look like a YouTube video link.")
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
    except ImportError:
        raise IngestError("YouTube support needs: pip install youtube-transcript-api") from None
    try:
        api = YouTubeTranscriptApi()
        fetched = api.fetch(vid) if hasattr(api, "fetch") else YouTubeTranscriptApi.get_transcript(vid)
        segments = [
            {"start": getattr(s, "start", None) if not isinstance(s, dict) else s.get("start"),
             "duration": getattr(s, "duration", 0) if not isinstance(s, dict) else s.get("duration", 0),
             "text": getattr(s, "text", "") if not isinstance(s, dict) else s.get("text", "")}
            for s in fetched
        ]
    except Exception:  # noqa: BLE001
        raise IngestError("No captions are available for that video (or YouTube blocked the request).") from None
    pages: list[Page] = []
    window_start, buf = 0.0, []
    end_time = 0.0
    for seg in segments:
        start = float(seg["start"] or 0)
        end_time = max(end_time, start + float(seg["duration"] or 0))
        if not buf:
            window_start = start
        buf.append(str(seg["text"]).replace("\n", " "))
        if start - window_start >= 60:
            pages.append(Page(None, " ".join(buf), _mmss(window_start)))
            buf = []
    if buf:
        pages.append(Page(None, " ".join(buf), _mmss(window_start)))
    if not pages:
        raise IngestError("The transcript is empty.")
    return pages, int(end_time)
